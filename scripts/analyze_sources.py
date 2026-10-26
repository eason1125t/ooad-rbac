"""只读审计根目录三个 Excel 来源，为 docs/12-课程交付与资料边界.adoc 提供可复现聚合事实。

仅输出表头、计数与权限词汇；不输出真实姓名、电话、账号或人员记录。
"""

from __future__ import annotations

import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import re
import unicodedata

import openpyxl


def normalize(value: object) -> str:
    """仅消除 Unicode 宽度与边缘空白，不猜测组织同义名称。"""
    return "" if value is None else unicodedata.normalize("NFKC", str(value)).strip()


def duplicate_stats(values: list[str]) -> dict[str, int]:
    counts = Counter(value for value in values if value)
    return {
        "distinct": len(counts),
        "duplicate_groups": sum(count > 1 for count in counts.values()),
        "excess_rows": sum(count - 1 for count in counts.values()),
    }


def base_metadata(path: Path, sheet: object, header_row: int) -> dict:
    records = list(sheet.iter_rows(min_row=header_row + 1, values_only=True))
    records = [row for row in records if any(normalize(value) for value in row)]
    headers = [normalize(cell.value) for cell in sheet[header_row]]
    return {
        "filename": path.name,
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "sheet": sheet.title,
        "header_row": header_row,
        "headers": headers,
        "records": len(records),
        "merged_ranges": len(sheet.merged_cells.ranges),
        "empty_by_column": {
            header: sum(not normalize(row[index]) for row in records)
            for index, header in enumerate(headers)
        },
        "normalization_changed_cells": sum(
            value is not None and normalize(value) != str(value)
            for row in records for value in row
        ),
        "formula_cells": sum(
            cell.data_type == "f" for row in sheet for cell in row
        ),
    }


def analyze(root: Path) -> dict:
    """合并仅用于读取 A/B 层级，权限列的空格不继承、不授予。"""
    paths = list(root.glob("*.xlsx"))
    employee_path = next(path for path in paths if "员工信息" in path.name)
    organization_path = next(path for path in paths if "部门信息" in path.name)
    permission_path = next(path for path in paths if "权限清单" in path.name)
    employee_book = openpyxl.load_workbook(employee_path, data_only=False)
    organization_book = openpyxl.load_workbook(organization_path, data_only=False)
    permission_book = openpyxl.load_workbook(permission_path, data_only=False)
    try:
        employees = employee_book.active
        organizations = organization_book.active
        permissions = permission_book.active
        employee_rows = [tuple(normalize(value) for value in row)
                         for row in employees.iter_rows(min_row=2, values_only=True)
                         if any(normalize(value) for value in row)]
        organization_rows = [tuple(normalize(value) for value in row)
                            for row in organizations.iter_rows(min_row=2, values_only=True)
                            if any(normalize(value) for value in row)]
        org_names = {row[0] for row in organization_rows}
        parents = {row[0]: row[2] for row in organization_rows}
        names = Counter(row[0] for row in employee_rows)
        prefix_matches = exact_matches = direct_branches = unresolved = ambiguous = 0
        branches = {row[2] for row in employee_rows}
        # 精确名称优先，再在已确认父节点中匹配唯一后缀，禁止盲目拼接前缀。
        for row in employee_rows:
            branch, department = row[2:4]
            if department == branch and department in org_names:
                direct_branches += 1
            elif department in org_names and parents[department] == branch:
                exact_matches += 1
            else:
                candidates = [name for name in org_names
                              if parents[name] == branch
                              and name in {branch + department, branch + "-" + department}]
                if len(candidates) == 1:
                    prefix_matches += 1
                elif len(candidates) > 1:
                    ambiguous += 1
                else:
                    unresolved += 1

        cycles = 0
        depths = Counter()
        for name in org_names:
            seen = set()
            current = name
            while current in parents:
                if current in seen:
                    cycles += 1
                    break
                seen.add(current)
                current = parents[current]
            if current not in parents:
                depths[len(seen) - 1] += 1

        org_meta = base_metadata(organization_path, organizations, 1)
        org_meta.update({
            "organization_names": duplicate_stats([row[0] for row in organization_rows]),
            "parent_labels_distinct": len({row[2] for row in organization_rows}),
            "nonempty_parent_not_in_nodes": len({row[2] for row in organization_rows if row[2] and row[2] not in org_names}),
            "parent_empty_rows": sum(not row[2] for row in organization_rows),
            "parent_reference_missing_rows": sum(bool(row[2]) and row[2] not in org_names for row in organization_rows),
            "cycle_affected_nodes": cycles,
            "nodes_by_depth": dict(sorted(depths.items())),
            "branch_labels_in_employees": len(branches),
            "branches_found_in_organizations": len(branches & org_names),
            "manager_distinct_names": len({row[1] for row in organization_rows if row[1]}),
            "manager_name_employee_match": dict(Counter(
                "absent" if names[row[1]] == 0 else "unique" if names[row[1]] == 1 else "ambiguous"
                for row in organization_rows if row[1]
            )),
        })
        employee_meta = base_metadata(employee_path, employees, 1)
        employee_meta.update({
            "usernames": duplicate_stats([row[1] for row in employee_rows]),
            "names": duplicate_stats([row[0] for row in employee_rows]),
            "phones": duplicate_stats([row[5] for row in employee_rows]),
            "phone_non_11_digit_rows": sum(not re.fullmatch(r"\d{11}", row[5]) for row in employee_rows),
            "job_labels_distinct": len({row[4] for row in employee_rows}),
            "employee_department_labels_distinct": len({row[3] for row in employee_rows}),
            "employee_organization_matches": {
                "direct_branch": direct_branches, "exact_child_with_parent": exact_matches,
                "unique_prefixed_child_with_parent": prefix_matches,
                "ambiguous": ambiguous, "unresolved": unresolved,
            },
        })

        effective = {}
        for merged in permissions.merged_cells.ranges:
            if merged.min_col in (1, 2) and merged.max_col == merged.min_col:
                value = normalize(permissions.cell(merged.min_row, merged.min_col).value)
                for row_number in range(max(3, merged.min_row), merged.max_row + 1):
                    effective[(row_number, merged.min_col)] = value
        permission_values = Counter()
        application_labels = set()
        module_keys = set()
        function_keys = []
        role_counts = {}
        special_permission_cells = []
        for row_number in range(3, permissions.max_row + 1):
            system = effective.get((row_number, 1), normalize(permissions.cell(row_number, 1).value))
            module = effective.get((row_number, 2), normalize(permissions.cell(row_number, 2).value))
            function = normalize(permissions.cell(row_number, 3).value)
            application_labels.add(system)
            module_keys.add((system, module))
            function_keys.append((system, module, function))
            permission_values.update(normalize(permissions.cell(row_number, col).value)
                                     for col in range(5, 10))
            for col in range(5, 10):
                value = normalize(permissions.cell(row_number, col).value)
                if not value or "本人" in value or "view" in value:
                    special_permission_cells.append({
                        "cell": permissions.cell(row_number, col).coordinate,
                        "application_label": system, "module_label": module,
                        "function_label": function,
                        "role_label": normalize(permissions.cell(2, col).value),
                        "value": value,
                    })
        for col in range(5, 10):
            values = Counter(normalize(permissions.cell(row, col).value)
                             for row in range(3, permissions.max_row + 1))
            role_counts[normalize(permissions.cell(2, col).value)] = {
                "none": values["无"], "blank": values[""], "all": values["全部"],
                "explicit_action_cells": sum(count for value, count in values.items()
                                             if value not in ("", "无", "全部")),
            }
        permission_meta = base_metadata(permission_path, permissions, 2)
        permission_meta.update({
            "applications": len(application_labels),
            "application_modules": len(module_keys),
            "duplicate_function_keys": len(function_keys) - len(set(function_keys)),
            "role_cells": sum(permission_values.values()),
            "permission_values": dict(permission_values),
            "role_counts": role_counts,
            "self_scope_cells": sum(count for value, count in permission_values.items() if "本人" in value),
            "english_view_cells": sum(count for value, count in permission_values.items() if "view" in value),
            "special_permission_cells": special_permission_cells,
        })
        screenshot_files = [path for path in root.iterdir() if path.is_file()
                            and path.suffix.lower() in {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".pdf"}]
        return {"organizations": org_meta, "employees": employee_meta,
                "permissions": permission_meta, "root_screenshot_or_pdf_count": len(screenshot_files)}
    finally:
        employee_book.close()
        organization_book.close()
        permission_book.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parent.parent)
    arguments = parser.parse_args()
    # ASCII 转义保证 Windows shell 的编码差异不影响 JSON，同时保持内容仅为聚合数据。
    print(json.dumps(analyze(arguments.root), ensure_ascii=True, indent=2))


if __name__ == "__main__":
    main()
